from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_processor import (
    SortWarningConfirmationRequired,
    UnsupportedWorkbookError,
    sort_xlsx,
)


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Names"
    worksheet.append(["Tên", "Mã", "Ghi chú"])
    worksheet.append(["Nguyễn Đạt", 3, "row-3"])
    worksheet.append(["Nguyễn An", 1, "row-2"])
    worksheet.append(["Nguyễn Đăng", 2, "row-4"])
    worksheet["A2"].font = Font(bold=True, color="FF0000")
    worksheet["C3"].comment = Comment("preserve me", "test")
    worksheet.row_dimensions[3].height = 30
    workbook.save(path)
    workbook.close()


def test_sorts_whole_rows_preserves_data_styles_comments_and_source(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    _make_workbook(source)

    result = sort_xlsx(
        source,
        output,
        worksheet_name="Names",
        sort_column="A",
        first_data_row=2,
        header_row=1,
        backend="fallback",
    )

    assert result.sorted_row_count == 3
    assert result.backend == "fallback"
    sorted_book = load_workbook(output)
    sorted_sheet = sorted_book["Names"]
    assert [sorted_sheet.cell(row, 1).value for row in range(2, 5)] == [
        "Nguyễn An",
        "Nguyễn Đạt",
        "Nguyễn Đăng",
    ]
    assert [sorted_sheet.cell(row, 2).value for row in range(2, 5)] == [1, 3, 2]
    assert [sorted_sheet.cell(row, 3).value for row in range(2, 5)] == [
        "row-2",
        "row-3",
        "row-4",
    ]
    assert sorted_sheet["A3"].font.bold
    assert sorted_sheet["A3"].font.color.rgb == "00FF0000"
    assert sorted_sheet["C2"].comment.text == "preserve me"
    assert sorted_sheet.row_dimensions[2].height == 30
    sorted_book.close()

    original = load_workbook(source)
    assert [original["Names"].cell(row, 1).value for row in range(2, 5)] == [
        "Nguyễn Đạt",
        "Nguyễn An",
        "Nguyễn Đăng",
    ]
    original.close()


def test_descending_keeps_blanks_at_selected_end(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "marker"])
    worksheet.append(["B", 1])
    worksheet.append([None, 2])
    worksheet.append(["A", 3])
    workbook.save(source)
    workbook.close()

    sort_xlsx(
        source,
        output,
        worksheet_name="Sheet",
        sort_column=1,
        first_data_row=2,
        descending=True,
        blanks_last=True,
        backend="fallback",
    )
    result = load_workbook(output, read_only=True)
    assert [result["Sheet"].cell(row, 1).value for row in range(2, 5)] == [
        "B",
        "A",
        None,
    ]
    assert [result["Sheet"].cell(row, 2).value for row in range(2, 5)] == [1, 3, 2]
    result.close()


def test_mixed_text_numbers_dates_and_blanks_have_defined_groups(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["value"])
    worksheet.append([date(2024, 1, 1)])
    worksheet.append([3])
    worksheet.append(["An"])
    worksheet.append([None])
    worksheet.append([1])
    workbook.save(source)
    workbook.close()

    sort_xlsx(
        source,
        output,
        worksheet_name="Sheet",
        sort_column=1,
        first_data_row=2,
        backend="fallback",
    )
    result = load_workbook(output, read_only=True)
    assert [result["Sheet"].cell(row, 1).value for row in range(2, 7)] == [
        "An",
        1,
        3,
        datetime(2024, 1, 1),
        None,
    ]
    result.close()


def test_formula_in_any_moved_cell_blocks_sort(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "calculated"])
    worksheet.append(["B", "=1+1"])
    worksheet.append(["A", 3])
    workbook.save(source)
    workbook.close()

    with pytest.raises(UnsupportedWorkbookError, match="Formulas occur"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=2,
            backend="fallback",
        )
    assert not output.exists()


def test_formula_outside_moved_rows_requires_acknowledgement(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name"])
    worksheet.append(["B"])
    worksheet.append(["A"])
    worksheet["A5"] = "=COUNTA(A2:A3)"
    workbook.save(source)
    workbook.close()

    with pytest.raises(SortWarningConfirmationRequired, match="Formulas outside"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=2,
            last_data_row=3,
            backend="fallback",
        )
    result = sort_xlsx(
        source,
        output,
        worksheet_name="Sheet",
        sort_column=1,
        first_data_row=2,
        last_data_row=3,
        backend="fallback",
        acknowledge_warnings=True,
    )
    assert result.warnings
    assert output.exists()


def test_formula_on_another_sheet_requires_acknowledgement(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name"])
    worksheet.append(["B"])
    worksheet.append(["A"])
    workbook.create_sheet("Summary")["A1"] = "=COUNTA(Sheet!A2:A3)"
    workbook.save(source)
    workbook.close()

    with pytest.raises(SortWarningConfirmationRequired, match="Summary!A1"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=2,
            last_data_row=3,
            backend="fallback",
        )


@pytest.mark.parametrize(
    ("first_value", "expected_message"),
    [
        (True, "Boolean values"),
        ("#N/A", "Excel error value"),
    ],
)
def test_unsupported_sort_values_are_rejected(
    tmp_path: Path,
    first_value: object,
    expected_message: str,
) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["value"])
    worksheet.append([first_value])
    worksheet.append(["An"])
    if first_value == "#N/A":
        worksheet["A2"].data_type = "e"
    workbook.save(source)
    workbook.close()

    with pytest.raises(UnsupportedWorkbookError, match=expected_message):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=2,
            backend="fallback",
        )
    assert not output.exists()


def test_merged_cells_intersecting_rows_block_sort(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "other"])
    worksheet.append(["B", None])
    worksheet.append(["A", None])
    worksheet.merge_cells("B2:B3")
    workbook.save(source)
    workbook.close()

    with pytest.raises(UnsupportedWorkbookError, match="Merged cells"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=2,
            backend="fallback",
        )


def test_excel_table_intersecting_rows_blocks_sort(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name"])
    worksheet.append(["B"])
    worksheet.append(["A"])
    table = Table(displayName="NamesTable", ref="A1:A3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    worksheet.add_table(table)
    workbook.save(source)
    workbook.close()

    with pytest.raises(UnsupportedWorkbookError, match="Excel tables"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=2,
            backend="fallback",
        )


@pytest.mark.parametrize(
    ("source_name", "output_name", "error", "message"),
    [
        ("input.xlsm", "sorted.xlsx", ValueError, "Only .xlsx"),
        ("input.xlsx", "sorted.xlsm", ValueError, "output path must use"),
    ],
)
def test_rejects_unsupported_extensions(
    tmp_path: Path,
    source_name: str,
    output_name: str,
    error: type[Exception],
    message: str,
) -> None:
    source = tmp_path / source_name
    output = tmp_path / output_name
    source.write_bytes(b"not opened")

    with pytest.raises(error, match=message):
        sort_xlsx(
            source,
            output,
            worksheet_name="Sheet",
            sort_column=1,
            first_data_row=1,
        )


def test_never_overwrites_source_or_existing_output(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    _make_workbook(source)
    original_bytes = source.read_bytes()
    with pytest.raises(ValueError, match="must differ"):
        sort_xlsx(
            source,
            source,
            worksheet_name="Names",
            sort_column=1,
            first_data_row=2,
        )
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"keep")
    with pytest.raises(FileExistsError, match="will not be overwritten"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Names",
            sort_column=1,
            first_data_row=2,
        )
    assert source.read_bytes() == original_bytes
    assert output.read_bytes() == b"keep"


def test_invalid_sheet_range_and_sort_column_are_reported(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    _make_workbook(source)

    with pytest.raises(ValueError, match="not found"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Missing",
            sort_column=1,
            first_data_row=2,
        )
    with pytest.raises(ValueError, match="below the optional header_row"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Names",
            sort_column=1,
            first_data_row=1,
            header_row=1,
        )
    with pytest.raises(ValueError, match="exceeds worksheet last column"):
        sort_xlsx(
            source,
            output,
            worksheet_name="Names",
            sort_column="D",
            first_data_row=2,
        )


def test_nfc_and_nfd_values_sort_equal_without_rewriting_originals(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "sorted.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name"])
    worksheet.append(["Nguye\u0302\u0303n"])
    worksheet.append(["Nguyễn"])
    workbook.save(source)
    workbook.close()

    sort_xlsx(
        source,
        output,
        worksheet_name="Sheet",
        sort_column=1,
        first_data_row=2,
        backend="fallback",
    )
    result = load_workbook(output, read_only=True)
    actual = [result["Sheet"].cell(row, 1).value for row in range(2, 4)]
    assert set(actual) == {"Nguye\u0302\u0303n", "Nguyễn"}
    assert actual[0] == "Nguye\u0302\u0303n"
    result.close()
