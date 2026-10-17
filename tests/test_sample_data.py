import csv
from pathlib import Path

from openpyxl import load_workbook

from excel_processor import sort_xlsx


def test_sample_workbook_matches_published_fallback_expected_order(tmp_path: Path) -> None:
    project_root = Path(__file__).resolve().parents[1]
    sample_path = project_root / "examples" / "sample_vietnamese_names.xlsx"
    expected_path = project_root / "examples" / "expected_vietnamese_names.csv"
    output_path = tmp_path / "sample_sorted.xlsx"

    with expected_path.open(encoding="utf-8", newline="") as source:
        expected = [
            (int(record_id), name, place)
            for record_id, name, place in list(csv.reader(source))[1:]
        ]

    sort_xlsx(
        sample_path,
        output_path,
        worksheet_name="Dữ liệu",
        sort_column="B",
        first_data_row=2,
        last_data_row=14,
        header_row=1,
        backend="fallback",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=False)
    try:
        worksheet = workbook["Dữ liệu"]
        actual = [
            tuple(worksheet.cell(row, column).value for column in range(1, 4))
            for row in range(2, 15)
        ]
    finally:
        workbook.close()
    assert actual == expected
