"""Safe `.xlsx` row sorting using the Vietnamese sorter."""

from __future__ import annotations

import os
import warnings
from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from vietnamese_sort import Backend, NameMode, VietnameseSorter

class ExcelProcessingError(Exception):
    """Base exception for workbook validation and sort errors."""


class UnsupportedWorkbookError(ExcelProcessingError):
    """Raised when sorting could silently damage unsupported workbook data."""


class SortWarningConfirmationRequired(ExcelProcessingError):
    """Raised when workbook warnings need explicit confirmation before saving."""

    def __init__(self, warnings_found: Sequence[str]) -> None:
        self.warnings_found = tuple(warnings_found)
        message = "Workbook warnings require acknowledgement: " + "; ".join(
            self.warnings_found
        )
        super().__init__(message)


@dataclass(frozen=True)
class SortResult:
    output_path: Path
    worksheet: str
    sort_column: int
    first_data_row: int
    last_data_row: int
    sorted_row_count: int
    warnings: tuple[str, ...]
    backend: str
    backend_locale: str | None


@dataclass(frozen=True)
class _CellSnapshot:
    value: object
    data_type: str
    style: object
    hyperlink: object
    comment: object


@dataclass(frozen=True)
class _RowSnapshot:
    cells: tuple[_CellSnapshot | None, ...]
    dimension: object | None


def _resolve_column(column: int | str) -> int:
    if isinstance(column, bool):
        raise ValueError("sort_column must be a 1-based number or Excel column label")
    if isinstance(column, int):
        index = column
    elif isinstance(column, str):
        try:
            index = column_index_from_string(column.strip().upper())
        except ValueError as error:
            raise ValueError(f"Invalid Excel column label: {column!r}") from error
    else:
        raise TypeError("sort_column must be a 1-based number or Excel column label")
    if index < 1:
        raise ValueError("sort_column must be greater than zero")
    return index


def _normalize_path(path: str | os.PathLike[str]) -> Path:
    return Path(path).expanduser().resolve()


def _validate_paths(
    source_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
) -> tuple[Path, Path]:
    source = _normalize_path(source_path)
    output = _normalize_path(output_path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx input workbooks are supported; .xlsm is not supported.")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("The output path must use the .xlsx extension.")
    if source == output:
        raise ValueError("The output path must differ from the source workbook.")
    if not source.is_file():
        raise FileNotFoundError(f"Source workbook does not exist: {source}")
    if output.exists():
        raise FileExistsError(f"Output file already exists and will not be overwritten: {output}")
    if not output.parent.is_dir():
        raise FileNotFoundError(f"Output folder does not exist: {output.parent}")
    return source, output


def _excel_value_key(value: object, epoch: datetime) -> tuple[int, object]:
    if isinstance(value, str):
        return 0, value
    if isinstance(value, bool):
        raise UnsupportedWorkbookError("Boolean values in the sort column are unsupported.")
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if not number.is_finite():
            raise UnsupportedWorkbookError("Non-finite numbers in the sort column are unsupported.")
        return 1, number
    if isinstance(value, (datetime, date, time, timedelta)):
        return 1, Decimal(str(to_excel(value, epoch)))
    raise UnsupportedWorkbookError(
        f"Unsupported sort-column value type: {type(value).__name__}."
    )


def _snapshot_row(worksheet: Worksheet, row: int, max_column: int) -> _RowSnapshot:
    cells: list[_CellSnapshot | None] = []
    for column in range(1, max_column + 1):
        cell = worksheet._cells.get((row, column))
        if cell is None:
            cells.append(None)
        else:
            cells.append(
                _CellSnapshot(
                    value=cell.value,
                    data_type=cell.data_type,
                    style=copy(cell._style),
                    hyperlink=copy(cell.hyperlink),
                    comment=copy(cell.comment),
                )
            )
    dimension = worksheet.row_dimensions.get(row)
    return _RowSnapshot(
        cells=tuple(cells),
        dimension=copy(dimension) if dimension is not None else None,
    )


def _write_row_snapshot(
    worksheet: Worksheet,
    row: int,
    snapshot: _RowSnapshot,
) -> None:
    for column, saved in enumerate(snapshot.cells, start=1):
        cell = worksheet.cell(row=row, column=column)
        cell.value = saved.value if saved is not None else None
        if saved is None:
            cell._style = None
            cell.hyperlink = None
            cell.comment = None
            continue
        cell.data_type = saved.data_type
        cell._style = copy(saved.style)
        cell.hyperlink = copy(saved.hyperlink)
        if cell.hyperlink is not None:
            cell.hyperlink.ref = cell.coordinate
        cell.comment = copy(saved.comment)

    worksheet.row_dimensions.pop(row, None)
    if snapshot.dimension is not None:
        dimension = copy(snapshot.dimension)
        dimension.index = row
        dimension.worksheet = worksheet
        worksheet.row_dimensions[row] = dimension


def _verify_saved_rows(
    worksheet: Worksheet,
    first_row: int,
    sorted_rows: Sequence[tuple[int, tuple]],
    snapshots: dict[int, _RowSnapshot],
) -> None:
    for destination_row, (source_row, _) in zip(
        range(first_row, first_row + len(sorted_rows)),
        sorted_rows,
        strict=True,
    ):
        expected_row = snapshots[source_row]
        for column, expected in enumerate(expected_row.cells, start=1):
            actual = worksheet.cell(row=destination_row, column=column)
            if expected is None:
                if actual.value is not None:
                    raise ExcelProcessingError(
                        f"Saved row verification failed at {actual.coordinate}."
                    )
                continue
            if (
                actual.value != expected.value
                or actual.data_type != expected.data_type
                or copy(actual._style) != expected.style
            ):
                raise ExcelProcessingError(
                    f"Saved row verification failed at {actual.coordinate}."
                )
            expected_link = expected.hyperlink
            if (actual.hyperlink.target if actual.hyperlink else None) != (
                expected_link.target if expected_link else None
            ):
                raise ExcelProcessingError(
                    f"Saved hyperlink verification failed at {actual.coordinate}."
                )
            expected_comment = expected.comment
            if (actual.comment.text if actual.comment else None) != (
                expected_comment.text if expected_comment else None
            ):
                raise ExcelProcessingError(
                    f"Saved comment verification failed at {actual.coordinate}."
                )


def _find_risks(
    workbook: Workbook,
    worksheet: Worksheet,
    first_row: int,
    last_row: int,
) -> tuple[list[str], list[str]]:
    blocking: list[str] = []
    warnings_found: list[str] = []

    if getattr(workbook, "_external_links", ()):
        blocking.append("Workbook contains external links.")

    merged_ranges = [
        str(merged)
        for merged in worksheet.merged_cells.ranges
        if merged.min_row <= last_row and merged.max_row >= first_row
    ]
    if merged_ranges:
        blocking.append(
            "Merged cells intersect rows being sorted: " + ", ".join(merged_ranges)
        )

    tables = []
    for table in worksheet.tables.values():
        if table.ref:
            _, min_row, _, max_row = range_boundaries(table.ref)
            if min_row <= last_row and max_row >= first_row:
                tables.append(table.displayName)
    if tables:
        blocking.append("Excel tables intersect rows being sorted: " + ", ".join(tables))

    if getattr(worksheet, "_charts", None):
        blocking.append("Worksheet contains charts that cannot be safely moved with rows.")
    if getattr(worksheet, "_images", None):
        blocking.append("Worksheet contains images that cannot be safely moved with rows.")
    if getattr(worksheet, "_pivots", None):
        blocking.append("Worksheet contains pivot tables that cannot be safely moved.")
    if worksheet.conditional_formatting:
        blocking.append(
            "Worksheet contains conditional formatting that does not move with sorted rows."
        )
    if worksheet.data_validations.dataValidation:
        blocking.append(
            "Worksheet contains data validations that do not move with sorted rows."
        )

    moved_formula_cells = [
        cell.coordinate
        for cell in worksheet._cells.values()
        if first_row <= cell.row <= last_row and cell.data_type == "f"
    ]
    if moved_formula_cells:
        blocking.append(
            "Formulas occur in rows being moved: " + ", ".join(moved_formula_cells)
        )

    if workbook.defined_names:
        warnings_found.append("Workbook contains defined names that may refer to moved rows.")

    outside_formula_cells = []
    for sheet in workbook.worksheets:
        for cell in sheet._cells.values():
            if cell.data_type == "f" and (
                sheet != worksheet or not first_row <= cell.row <= last_row
            ):
                outside_formula_cells.append(f"{sheet.title}!{cell.coordinate}")
    if outside_formula_cells:
        warnings_found.append(
            "Formulas outside the moved rows may refer to sorted cells: "
            + ", ".join(outside_formula_cells)
        )

    return blocking, warnings_found


def sort_xlsx(
    source_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
    *,
    worksheet_name: str,
    sort_column: int | str,
    first_data_row: int,
    last_data_row: int | None = None,
    header_row: int | None = None,
    descending: bool = False,
    blanks_last: bool = True,
    name_mode: NameMode = "full_name",
    backend: Backend = "auto",
    acknowledge_warnings: bool = False,
) -> SortResult:
    """Sort whole worksheet rows by one column and save to a new `.xlsx` file."""
    source, output = _validate_paths(source_path, output_path)
    column_index = _resolve_column(sort_column)

    if isinstance(first_data_row, bool) or first_data_row < 1:
        raise ValueError("first_data_row must be a positive 1-based row number.")
    if header_row is not None and (
        isinstance(header_row, bool)
        or header_row < 1
        or first_data_row <= header_row
    ):
        raise ValueError("first_data_row must be below the optional header_row.")
    if last_data_row is not None and (
        isinstance(last_data_row, bool) or last_data_row < first_data_row
    ):
        raise ValueError("last_data_row must be greater than or equal to first_data_row.")

    load_warnings: list[str] = []
    with warnings.catch_warnings(record=True) as captured_warnings:
        warnings.simplefilter("always")
        workbook = load_workbook(source, data_only=False, keep_links=True)
    load_warnings.extend(
        f"openpyxl: {item.message}"
        for item in captured_warnings
        if any(
            phrase in str(item.message).lower()
            for phrase in ("not supported", "will be removed", "cannot be preserved")
        )
    )
    if load_warnings:
        raise UnsupportedWorkbookError(
            "Workbook contains features openpyxl warned it may not preserve: "
            + "; ".join(load_warnings)
        )

    if worksheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet {worksheet_name!r} not found. Available sheets: "
            + ", ".join(workbook.sheetnames)
        )
    worksheet = workbook[worksheet_name]
    final_row = last_data_row if last_data_row is not None else worksheet.max_row
    if first_data_row > worksheet.max_row or final_row > worksheet.max_row:
        raise ValueError(
            f"Data row range {first_data_row}:{final_row} exceeds worksheet "
            f"last row {worksheet.max_row}."
        )
    if column_index > worksheet.max_column:
        raise ValueError(
            f"Sort column {column_index} exceeds worksheet last column "
            f"{worksheet.max_column}."
        )

    blocking, warnings_found = _find_risks(
        workbook,
        worksheet,
        first_data_row,
        final_row,
    )
    if blocking:
        raise UnsupportedWorkbookError(" ".join(blocking))
    if warnings_found and not acknowledge_warnings:
        raise SortWarningConfirmationRequired(warnings_found)

    sorter = VietnameseSorter(backend=backend)
    rows: list[tuple[int, tuple]] = []
    for row in range(first_data_row, final_row + 1):
        sort_cell = worksheet.cell(row=row, column=column_index)
        if sort_cell.data_type == "e":
            raise UnsupportedWorkbookError(
                f"Excel error value in sort cell {sort_cell.coordinate} is unsupported."
            )
        value = sort_cell.value
        if isinstance(value, str) and not value.strip():
            rows.append((row, (2, 0)))
        elif value is None:
            rows.append((row, (2, 0)))
        else:
            kind, sortable_value = _excel_value_key(value, workbook.epoch)
            if kind == 0:
                key = (0, sorter.key(sortable_value, name_mode))
            else:
                key = (1, sortable_value)
            rows.append((row, key))

    nonblank = [entry for entry in rows if entry[1][0] != 2]
    blanks = [entry for entry in rows if entry[1][0] == 2]
    nonblank.sort(key=lambda entry: entry[1], reverse=descending)
    sorted_rows = nonblank + blanks if blanks_last else blanks + nonblank

    snapshots = {
        row: _snapshot_row(worksheet, row, worksheet.max_column)
        for row in range(first_data_row, final_row + 1)
    }
    for destination_row, (source_row, _) in zip(
        range(first_data_row, final_row + 1),
        sorted_rows,
        strict=True,
    ):
        _write_row_snapshot(worksheet, destination_row, snapshots[source_row])

    output_created = False
    try:
        with output.open("xb") as stream:
            output_created = True
            workbook.save(stream)
        verified = load_workbook(output, data_only=False, keep_links=True)
        try:
            if worksheet_name not in verified.sheetnames:
                raise ExcelProcessingError("Saved workbook is missing the selected worksheet.")
            verified_sheet = verified[worksheet_name]
            if verified_sheet.max_column < worksheet.max_column:
                raise ExcelProcessingError("Saved workbook has fewer columns than expected.")
            _verify_saved_rows(
                verified_sheet,
                first_data_row,
                sorted_rows,
                snapshots,
            )
        finally:
            verified.close()
    except Exception:
        workbook.close()
        if output_created:
            output.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()

    return SortResult(
        output_path=output,
        worksheet=worksheet_name,
        sort_column=column_index,
        first_data_row=first_data_row,
        last_data_row=final_row,
        sorted_row_count=len(sorted_rows),
        warnings=tuple(warnings_found),
        backend=sorter.backend,
        backend_locale=sorter.backend_locale,
    )
