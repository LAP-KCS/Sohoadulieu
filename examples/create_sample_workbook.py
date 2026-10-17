"""Create a small synthetic workbook for manually checking the application."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


RECORDS = [
    (113, "Nguyễn Thị", "Huế"),
    (104, "Nguyễn Đạt", "Đà Nẵng"),
    (101, "Nguyễn An", "Hà Nội"),
    (112, "Nguyễn Văn", "Cần Thơ"),
    (105, "Nguyễn Đăng", "Hải Phòng"),
    (102, "Nguyễn Ánh", "Hà Nội"),
    (110, "Nguyễn Nam", "Cần Thơ"),
    (103, "Nguyễn Bình", "TP. HCM"),
    (108, "Nguyễn Hương", "Huế"),
    (106, "Nguyễn Đức", "Đà Nẵng"),
    (111, "Nguyễn Ngọc", "Hải Phòng"),
    (109, "Nguyễn Minh", "TP. HCM"),
    (107, "Nguyễn Hùng", "Hà Nội"),
]

OUTPUT = Path(__file__).with_name("sample_vietnamese_names.xlsx")


def main() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Dữ liệu"
    worksheet.append(["Mã bản ghi", "Họ tên", "Địa danh"])
    for record in RECORDS:
        worksheet.append(record)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 16
    worksheet.column_dimensions["C"].width = 20

    workbook.save(OUTPUT)
    workbook.close()
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
