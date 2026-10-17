"""Tkinter desktop interface for the Vietnamese Excel sorter."""

from __future__ import annotations

import os
import shutil
import tempfile
import tkinter as tk
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from excel_processor import (
    SortResult,
    SortWarningConfirmationRequired,
    UnsupportedWorkbookError,
    sort_xlsx,
)
from vietnamese_sort import Backend, NameMode


@dataclass(frozen=True)
class _SortOptions:
    source: Path
    worksheet: str
    column: int
    header_row: int | None
    first_data_row: int
    last_data_row: int | None
    descending: bool
    blanks_last: bool
    name_mode: NameMode
    backend: Backend


def _format_preview_value(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


class VietnameseExcelSorterApp:
    """Desktop app with a preview-first, no-overwrite workflow."""

    PREVIEW_LIMIT = 50

    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Sắp xếp tiếng Việt trong Excel")
        self.root.geometry("1050x720")
        self.root.minsize(800, 560)

        self.source_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.column_var = tk.StringVar()
        self.header_var = tk.StringVar(value="1")
        self.first_row_var = tk.StringVar(value="2")
        self.last_row_var = tk.StringVar()
        self.direction_var = tk.StringVar(value="Tăng dần")
        self.name_mode_var = tk.StringVar(value="Toàn bộ họ tên")
        self.blank_position_var = tk.StringVar(value="Cuối")
        self.backend_var = tk.StringVar(value="Tự động")
        self.status_var = tk.StringVar(value="Chọn tệp .xlsx để bắt đầu.")
        self._column_by_label: dict[str, int] = {}
        self._preview_directory: tempfile.TemporaryDirectory[str] | None = None
        self._preview_path: Path | None = None
        self._preview_options: _SortOptions | None = None
        self._preview_signature: tuple[int, int] | None = None
        self._preview_result: SortResult | None = None
        self._last_output_path: Path | None = None

        self._build_widgets()
        self.root.protocol("WM_DELETE_WINDOW", self._close)

    def _build_widgets(self) -> None:
        container = ttk.Frame(self.root, padding=12)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(1, weight=1)
        container.rowconfigure(5, weight=1)

        ttk.Label(container, text="Tệp Excel (.xlsx):").grid(
            row=0, column=0, sticky=tk.W, padx=(0, 8), pady=4
        )
        ttk.Entry(container, textvariable=self.source_var).grid(
            row=0, column=1, sticky=tk.EW, pady=4
        )
        ttk.Button(container, text="Chọn tệp...", command=self._choose_file).grid(
            row=0, column=2, padx=(8, 0), pady=4
        )

        settings = ttk.LabelFrame(container, text="Thiết lập sắp xếp", padding=8)
        settings.grid(row=1, column=0, columnspan=3, sticky=tk.EW, pady=8)
        for column in (1, 3, 5):
            settings.columnconfigure(column, weight=1)

        self._add_combo(settings, "Trang tính:", self.sheet_var, row=0, column=0)
        self.sheet_combo = self._last_combo
        self.sheet_combo.bind("<<ComboboxSelected>>", self._select_sheet)
        self._add_combo(settings, "Cột khóa:", self.column_var, row=0, column=2)
        self.column_combo = self._last_combo

        self._add_entry(settings, "Hàng tiêu đề:", self.header_var, row=1, column=0)
        self._add_entry(settings, "Dòng dữ liệu từ:", self.first_row_var, row=1, column=2)
        self._add_entry(settings, "Đến (để trống = cuối):", self.last_row_var, row=1, column=4)

        self._add_combo(
            settings,
            "Thứ tự:",
            self.direction_var,
            row=2,
            column=0,
            values=("Tăng dần", "Giảm dần"),
        )
        self._add_combo(
            settings,
            "Sắp xếp tên:",
            self.name_mode_var,
            row=2,
            column=2,
            values=("Toàn bộ họ tên", "Tên gọi cuối cùng"),
        )
        self._add_combo(
            settings,
            "Ô trống:",
            self.blank_position_var,
            row=2,
            column=4,
            values=("Cuối", "Đầu"),
        )

        self._add_combo(
            settings,
            "Bộ so sánh:",
            self.backend_var,
            row=3,
            column=0,
            values=("Tự động", "Fallback tiếng Việt"),
        )
        ttk.Label(
            settings,
            text="Tệp nguồn luôn được giữ nguyên; công thức trong hàng sắp xếp sẽ bị chặn.",
            wraplength=600,
        ).grid(row=3, column=2, columnspan=4, sticky=tk.W, padx=6, pady=4)

        actions = ttk.Frame(container)
        actions.grid(row=2, column=0, columnspan=3, sticky=tk.EW, pady=(0, 8))
        self.preview_button = ttk.Button(
            actions, text="Xem trước kết quả", command=self._preview
        )
        self.preview_button.pack(side=tk.LEFT)
        self.save_button = ttk.Button(
            actions,
            text="Thực hiện và lưu thành tệp mới...",
            command=self._save_result,
            state=tk.DISABLED,
        )
        self.save_button.pack(side=tk.LEFT, padx=8)
        ttk.Button(actions, text="Mở thư mục kết quả", command=self._open_output_folder).pack(
            side=tk.LEFT
        )

        ttk.Label(container, text="Bản xem trước (tối đa 50 hàng):").grid(
            row=3, column=0, columnspan=3, sticky=tk.W, pady=(4, 2)
        )
        self.preview_frame = ttk.Frame(container)
        self.preview_frame.grid(row=4, column=0, columnspan=3, sticky=tk.NSEW)
        container.rowconfigure(4, weight=1)
        self.preview_tree = ttk.Treeview(self.preview_frame, show="headings", height=14)
        vertical = ttk.Scrollbar(
            self.preview_frame, orient=tk.VERTICAL, command=self.preview_tree.yview
        )
        horizontal = ttk.Scrollbar(
            self.preview_frame, orient=tk.HORIZONTAL, command=self.preview_tree.xview
        )
        self.preview_tree.configure(
            yscrollcommand=vertical.set,
            xscrollcommand=horizontal.set,
        )
        self.preview_tree.grid(row=0, column=0, sticky=tk.NSEW)
        vertical.grid(row=0, column=1, sticky=tk.NS)
        horizontal.grid(row=1, column=0, sticky=tk.EW)
        self.preview_frame.rowconfigure(0, weight=1)
        self.preview_frame.columnconfigure(0, weight=1)

        ttk.Label(container, textvariable=self.status_var).grid(
            row=5, column=0, columnspan=3, sticky=tk.W, pady=(6, 2)
        )
        ttk.Label(container, text="Nhật ký xử lý:").grid(
            row=6, column=0, columnspan=3, sticky=tk.W, pady=(4, 2)
        )
        self.log = scrolledtext.ScrolledText(container, height=7, state=tk.DISABLED)
        self.log.grid(row=7, column=0, columnspan=3, sticky=tk.EW)

        for variable in (
            self.column_var,
            self.header_var,
            self.first_row_var,
            self.last_row_var,
            self.direction_var,
            self.name_mode_var,
            self.blank_position_var,
            self.backend_var,
        ):
            variable.trace_add("write", self._configuration_changed)

    def _add_combo(
        self,
        parent: ttk.Widget,
        label: str,
        variable: tk.StringVar,
        *,
        row: int,
        column: int,
        values: tuple[str, ...] = (),
    ) -> None:
        ttk.Label(parent, text=label).grid(
            row=row, column=column, sticky=tk.W, padx=6, pady=4
        )
        combo = ttk.Combobox(
            parent,
            textvariable=variable,
            values=values,
            state="readonly",
            width=24,
        )
        combo.grid(row=row, column=column + 1, sticky=tk.EW, padx=6, pady=4)
        self._last_combo = combo

    @staticmethod
    def _add_entry(
        parent: ttk.Widget,
        label: str,
        variable: tk.StringVar,
        *,
        row: int,
        column: int,
    ) -> None:
        ttk.Label(parent, text=label).grid(
            row=row, column=column, sticky=tk.W, padx=6, pady=4
        )
        ttk.Entry(parent, textvariable=variable, width=12).grid(
            row=row, column=column + 1, sticky=tk.EW, padx=6, pady=4
        )

    def _log(self, message: str) -> None:
        self.log.configure(state=tk.NORMAL)
        self.log.insert(tk.END, message.rstrip() + "\n")
        self.log.see(tk.END)
        self.log.configure(state=tk.DISABLED)

    def _configuration_changed(self, *_: object) -> None:
        self._discard_preview()

    def _discard_preview(self) -> None:
        if self._preview_directory is not None:
            self._preview_directory.cleanup()
        self._preview_directory = None
        self._preview_path = None
        self._preview_options = None
        self._preview_signature = None
        self._preview_result = None
        self.save_button.configure(state=tk.DISABLED)
        for item in self.preview_tree.get_children():
            self.preview_tree.delete(item)

    def _choose_file(self) -> None:
        selected = filedialog.askopenfilename(
            title="Chọn tệp Excel",
            filetypes=(("Excel Workbook (*.xlsx)", "*.xlsx"),),
        )
        if not selected:
            return
        self._discard_preview()
        self.source_var.set(selected)
        self._load_workbook_details(Path(selected))

    def _load_workbook_details(self, source: Path) -> None:
        try:
            if source.suffix.lower() != ".xlsx":
                raise ValueError("Chỉ hỗ trợ tệp .xlsx; tệp .xlsm không được hỗ trợ.")
            workbook = load_workbook(source, read_only=True, data_only=False)
            try:
                sheet_names = workbook.sheetnames
                self.sheet_combo.configure(values=sheet_names)
                if not sheet_names:
                    raise ValueError("Tệp không có trang tính.")
                self.sheet_var.set(sheet_names[0])
                self._populate_column_choices(workbook[sheet_names[0]])
                self._populate_row_defaults(workbook[sheet_names[0]])
            finally:
                workbook.close()
            self.status_var.set(f"Đã nạp {source.name}. Chọn vùng rồi xem trước.")
            self._log(f"Đã mở danh sách trang tính: {source}")
        except Exception as error:
            self._show_error(error)

    def _populate_column_choices(self, worksheet) -> None:
        header_row = self._optional_row(self.header_var.get(), "hàng tiêu đề")
        labels: list[str] = []
        self._column_by_label.clear()
        for column in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column)
            header = (
                worksheet.cell(header_row, column).value
                if header_row is not None and header_row <= worksheet.max_row
                else None
            )
            label = f"{letter} — {header}" if header not in (None, "") else letter
            labels.append(label)
            self._column_by_label[label] = column
        self.column_combo.configure(values=labels)
        if labels:
            self.column_var.set(labels[0])

    def _populate_row_defaults(self, worksheet) -> None:
        self.header_var.set("1")
        self.first_row_var.set("2" if worksheet.max_row > 1 else "1")
        self.last_row_var.set(str(worksheet.max_row))
        self._populate_column_choices(worksheet)

    def _select_sheet(self, _event: tk.Event | None = None) -> None:
        source_text = self.source_var.get().strip()
        if not source_text:
            return
        try:
            workbook = load_workbook(source_text, read_only=True, data_only=False)
            try:
                worksheet = workbook[self.sheet_var.get()]
                self._populate_row_defaults(worksheet)
            finally:
                workbook.close()
            self._configuration_changed()
        except Exception as error:
            self._show_error(error)

    @staticmethod
    def _optional_row(value: str, label: str) -> int | None:
        if not value.strip():
            return None
        try:
            row = int(value.strip())
        except ValueError as error:
            raise ValueError(f"{label} phải là số nguyên dương.") from error
        if row < 1:
            raise ValueError(f"{label} phải là số nguyên dương.")
        return row

    def _current_options(self) -> _SortOptions:
        source_text = self.source_var.get().strip()
        if not source_text:
            raise ValueError("Hãy chọn tệp Excel trước.")
        label = self.column_var.get()
        if label not in self._column_by_label:
            raise ValueError("Hãy chọn cột cần sắp xếp.")
        header_row = self._optional_row(self.header_var.get(), "Hàng tiêu đề")
        first_data_row = self._optional_row(
            self.first_row_var.get(), "Dòng dữ liệu bắt đầu"
        )
        if first_data_row is None:
            raise ValueError("Dòng dữ liệu bắt đầu là bắt buộc.")
        last_data_row = self._optional_row(
            self.last_row_var.get(), "Dòng dữ liệu kết thúc"
        )
        return _SortOptions(
            source=Path(source_text).expanduser().resolve(),
            worksheet=self.sheet_var.get(),
            column=self._column_by_label[label],
            header_row=header_row,
            first_data_row=first_data_row,
            last_data_row=last_data_row,
            descending=self.direction_var.get() == "Giảm dần",
            blanks_last=self.blank_position_var.get() == "Cuối",
            name_mode=(
                "given_name"
                if self.name_mode_var.get() == "Tên gọi cuối cùng"
                else "full_name"
            ),
            backend=(
                "fallback"
                if self.backend_var.get() == "Fallback tiếng Việt"
                else "auto"
            ),
        )

    @staticmethod
    def _file_signature(path: Path) -> tuple[int, int]:
        metadata = path.stat()
        return metadata.st_size, metadata.st_mtime_ns

    def _preview(self) -> None:
        self._discard_preview()
        try:
            options = self._current_options()
            signature = self._file_signature(options.source)
            self._preview_directory = tempfile.TemporaryDirectory(
                prefix="vietnamese-excel-sort-"
            )
            preview_path = Path(self._preview_directory.name) / "preview.xlsx"
            acknowledged = False
            try:
                result = self._run_sort(options, preview_path, acknowledged)
            except SortWarningConfirmationRequired as warning:
                self._log(str(warning))
                if not messagebox.askyesno(
                    "Xác nhận cảnh báo",
                    str(warning) + "\n\nBạn có muốn tiếp tục xem trước không?",
                    parent=self.root,
                ):
                    self._discard_preview()
                    return
                acknowledged = True
                result = self._run_sort(options, preview_path, acknowledged)

            if self._file_signature(options.source) != signature:
                raise RuntimeError("Tệp nguồn đã thay đổi trong lúc tạo bản xem trước.")
            self._preview_path = preview_path
            self._preview_options = options
            self._preview_signature = signature
            self._preview_result = result
            self._display_preview(options, preview_path)
            self.save_button.configure(state=tk.NORMAL)
            self.status_var.set(
                f"Đã xem trước {result.sorted_row_count} hàng bằng {result.backend}."
            )
            self._log(
                f"Xem trước hoàn tất: {result.sorted_row_count} hàng; "
                f"backend={result.backend}; locale={result.backend_locale or 'chưa xác minh'}."
            )
            if result.warnings:
                self._log("Cảnh báo đã được người dùng xác nhận: " + "; ".join(result.warnings))
        except Exception as error:
            self._discard_preview()
            self._show_error(error)

    @staticmethod
    def _run_sort(
        options: _SortOptions,
        output: Path,
        acknowledged: bool,
    ) -> SortResult:
        return sort_xlsx(
            options.source,
            output,
            worksheet_name=options.worksheet,
            sort_column=options.column,
            first_data_row=options.first_data_row,
            last_data_row=options.last_data_row,
            header_row=options.header_row,
            descending=options.descending,
            blanks_last=options.blanks_last,
            name_mode=options.name_mode,
            backend=options.backend,
            acknowledge_warnings=acknowledged,
        )

    def _display_preview(self, options: _SortOptions, preview_path: Path) -> None:
        workbook = load_workbook(preview_path, read_only=True, data_only=False)
        try:
            worksheet = workbook[options.worksheet]
            headers: list[str] = []
            for column in range(1, worksheet.max_column + 1):
                value = (
                    worksheet.cell(options.header_row, column).value
                    if options.header_row is not None
                    else None
                )
                headers.append(
                    _format_preview_value(value) or get_column_letter(column)
                )
            columns = ("__row__",) + tuple(
                f"column_{column}" for column in range(1, worksheet.max_column + 1)
            )
            self.preview_tree.configure(columns=columns)
            self.preview_tree.heading("__row__", text="Hàng")
            self.preview_tree.column("__row__", width=60, stretch=False)
            for column, header in enumerate(headers, start=1):
                column_id = f"column_{column}"
                self.preview_tree.heading(column_id, text=header)
                self.preview_tree.column(column_id, width=150, stretch=True)
            for item in self.preview_tree.get_children():
                self.preview_tree.delete(item)
            last_row = (
                options.last_data_row
                if options.last_data_row is not None
                else worksheet.max_row
            )
            preview_end = min(last_row, options.first_data_row + self.PREVIEW_LIMIT - 1)
            for row in range(options.first_data_row, preview_end + 1):
                values = [
                    _format_preview_value(worksheet.cell(row, column).value)
                    for column in range(1, worksheet.max_column + 1)
                ]
                self.preview_tree.insert("", tk.END, values=(row, *values))
            self._log(
                f"Hiển thị {max(0, preview_end - options.first_data_row + 1)} "
                f"hàng đầu trong vùng kết quả."
            )
        finally:
            workbook.close()

    def _save_result(self) -> None:
        if (
            self._preview_path is None
            or self._preview_options is None
            or self._preview_signature is None
        ):
            messagebox.showinfo(
                "Chưa có bản xem trước",
                "Hãy tạo bản xem trước trước khi lưu.",
                parent=self.root,
            )
            return
        try:
            if self._current_options() != self._preview_options:
                raise RuntimeError("Thiết lập đã thay đổi; hãy tạo lại bản xem trước.")
            if self._file_signature(self._preview_options.source) != self._preview_signature:
                raise RuntimeError("Tệp nguồn đã thay đổi; hãy tạo lại bản xem trước.")
            output_text = filedialog.asksaveasfilename(
                title="Lưu kết quả thành tệp Excel mới",
                initialdir=str(self._preview_options.source.parent),
                initialfile=f"{self._preview_options.source.stem}_sorted.xlsx",
                defaultextension=".xlsx",
                filetypes=(("Excel Workbook (*.xlsx)", "*.xlsx"),),
            )
            if not output_text:
                return
            output = Path(output_text).expanduser().resolve()
            if output.suffix.lower() != ".xlsx":
                raise ValueError("Tệp kết quả phải có phần mở rộng .xlsx.")
            if output == self._preview_options.source:
                raise ValueError("Tệp kết quả phải khác tệp nguồn.")
            if output.exists():
                raise FileExistsError(
                    f"Tệp đã tồn tại và sẽ không bị ghi đè: {output}"
                )
            self._copy_new_file(self._preview_path, output)
            self._last_output_path = output
            self.status_var.set(f"Đã lưu kết quả: {output}")
            self._log(f"Đã lưu tệp kết quả mới: {output}")
            messagebox.showinfo(
                "Đã lưu kết quả",
                f"Đã lưu {self._preview_result.sorted_row_count} hàng vào:\n{output}",
                parent=self.root,
            )
        except Exception as error:
            self._show_error(error)

    @staticmethod
    def _copy_new_file(source: Path, output: Path) -> None:
        created = False
        try:
            with source.open("rb") as input_stream, output.open("xb") as output_stream:
                created = True
                shutil.copyfileobj(input_stream, output_stream)
                output_stream.flush()
                os.fsync(output_stream.fileno())
        except Exception:
            if created:
                output.unlink(missing_ok=True)
            raise

    def _open_output_folder(self) -> None:
        path: Path | None = None
        if self._last_output_path is not None:
            path = self._last_output_path.parent
        elif self.source_var.get().strip():
            path = Path(self.source_var.get()).expanduser().resolve().parent
        if path is None:
            messagebox.showinfo(
                "Chưa chọn thư mục",
                "Hãy chọn tệp Excel trước.",
                parent=self.root,
            )
            return
        try:
            os.startfile(path)  # type: ignore[attr-defined]
        except (AttributeError, OSError) as error:
            self._show_error(RuntimeError(f"Không mở được thư mục {path}: {error}"))

    def _show_error(self, error: Exception) -> None:
        message = str(error) or type(error).__name__
        self.status_var.set("Đã xảy ra lỗi.")
        self._log(f"LỖI: {message}")
        messagebox.showerror("Không thể xử lý tệp Excel", message, parent=self.root)

    def _close(self) -> None:
        self._discard_preview()
        self.root.destroy()


def main() -> None:
    root = tk.Tk()
    VietnameseExcelSorterApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
